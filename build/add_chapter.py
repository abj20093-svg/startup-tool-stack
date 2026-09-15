#!/usr/bin/env python3
"""Append a chapter block to a copy of the source workbook.

Reads startup-tool-stack.xlsx (never written to), appends the block described
in a chapter JSON file after the last used row (one blank row between), and
writes startup-tool-stack-v2.xlsx. Cell styles and row heights are copied
from the matching rows of the last existing chapter so the block looks like
the others in Excel. Then reloads both files and asserts every original cell
(A1:Y1000) is unchanged.

Chapter JSON shape:
  {"group": "OPERATING", "name": "Compliance Automation",
   "tools": [6 vendor names..., "Manual DIY"],
   "rows": [{"label": "Built for", "cells": [7 strings]}, ...]}

usage: /usr/bin/python3 add_chapter.py compliance_chapter.json
"""
import copy, json, sys, os
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "startup-tool-stack.xlsx")
DST = os.path.join(HERE, "startup-tool-stack-v2.xlsx")
ch = json.load(open(sys.argv[1], encoding="utf-8"))
assert len(ch["tools"]) == 7 and ch["tools"][-1] == "Manual DIY", ch["tools"]
for r in ch["rows"]:
    assert len(r["cells"]) == 7, r["label"]
    assert all(isinstance(c, str) and c.strip() for c in r["cells"]), f"empty cell in {r['label']}"

wb = openpyxl.load_workbook(SRC)
ws = wb.worksheets[0]
# last used row
last = max(r for r in range(1, ws.max_row + 1) if any(ws.cell(row=r, column=c).value not in (None, "") for c in range(1, 10)))
header_row = last + 2                      # one blank row between chapters, as elsewhere
# template rows from the last existing chapter (Workflow Automation: header 200, first 201, choose 211)
T_HEADER, T_BODY, T_CHOOSE = 200, 201, 211
def copy_style(src_r, dst_r):
    for c in range(1, 10):
        ws.cell(row=dst_r, column=c)._style = copy.copy(ws.cell(row=src_r, column=c)._style)
    if ws.row_dimensions[src_r].height:
        ws.row_dimensions[dst_r].height = ws.row_dimensions[src_r].height

ws.cell(row=header_row, column=1, value="Sub-Niche")
ws.cell(row=header_row, column=2, value="Factor")
for i, t in enumerate(ch["tools"]):
    ws.cell(row=header_row, column=3 + i, value=t)
copy_style(T_HEADER, header_row)
r = header_row
for row in ch["rows"]:
    r += 1
    ws.cell(row=r, column=1, value=ch["name"])
    ws.cell(row=r, column=2, value=row["label"])
    for i, t in enumerate(row["cells"]):
        ws.cell(row=r, column=3 + i, value=t)
    copy_style(T_CHOOSE if row["label"].startswith("Choose this") else T_BODY, r)
first_row, last_row = header_row + 1, r
wb.save(DST)

# verify: every original cell unchanged, new block present
orig = openpyxl.load_workbook(SRC).worksheets[0]
new = openpyxl.load_workbook(DST).worksheets[0]
diffs = [(rr, cc) for rr in range(1, 1001) for cc in range(1, 26)
         if rr < header_row and orig.cell(row=rr, column=cc).value != new.cell(row=rr, column=cc).value]
assert not diffs, diffs[:10]
for rr in range(header_row + 1, last_row + 1):
    assert all(new.cell(row=rr, column=cc).value not in (None, "") for cc in range(1, 10)), rr
print(json.dumps({"dst": DST, "header_row": header_row, "first_row": first_row, "last_row": last_row,
                  "original_cells_changed": len(diffs)}, indent=1))
